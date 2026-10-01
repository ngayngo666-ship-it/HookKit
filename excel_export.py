#!/usr/bin/env python3
"""Xuat so quy ra file Excel (.xlsx) chuyen nghiep: Tong hop + moi ky chot mot sheet + ky dang mo.

    from ledger import Ledger
    from excel_export import export_xlsx
    export_xlsx(Ledger("ket.db"), "nhom-ket", "chot_so.xlsx")

- Cac so tong trong moi sheet la CONG THUC Excel tinh tu cac dong bill, ai mo file cung kiem tra lai duoc.
- O "Kiem tra khop so" so sanh voi so du cuoi ky ghi trong so quy (SQLite): KHOP / LECH.
- File chi la ban bao cao. So lieu goc luon nam trong SQLite, khong doc nguoc tu Excel.
"""
from __future__ import annotations

import time

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

NUM = '#,##0;[Red]-#,##0;0'
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
SUB_FILL = PatternFill("solid", fgColor="D9E2F3")
TOTAL_FILL = PatternFill("solid", fgColor="F2F2F2")
GREEN = PatternFill("solid", bgColor="C6EFCE", fgColor="C6EFCE")
RED = PatternFill("solid", bgColor="FFC7CE", fgColor="FFC7CE")
AMBER = PatternFill("solid", bgColor="FFEB9C", fgColor="FFEB9C")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

COLS = [("STT", 6), ("Thời gian", 17), ("Mã bill", 9), ("Mã GD sổ", 10), ("Loại", 14), ("Người", 20), ("STK", 16),
        ("Ngân hàng", 16), ("Nội dung", 34), ("Số tiền (không phí)", 20), ("Phí", 14), ("Số còn lại (sau phí)", 20),
        ("Người thao tác", 14)]
FIRST_DATA_ROW = 20
SUMMARY_ROWS = {
    "opening": 5, "total_in": 6, "total_out": 7, "fee_in": 8, "fee_out": 9, "fee_income": 10,
    "net_in": 11, "net_out": 12, "closing": 13, "position": 14, "check": 15,
}


def _fmt_time(ts) -> str:
    return time.strftime("%d/%m/%Y %H:%M", time.localtime(ts)) if ts else ""


def _title(ws, text, sub=None):
    ws["A1"] = text
    ws["A1"].font = Font(bold=True, size=15, color="1F3864")
    if sub:
        ws["A2"] = sub
        ws["A2"].font = Font(italic=True, color="595959")


def _period_sheet(wb, name, title, report, subtitle, ledger_closing):
    ws = wb.create_sheet(name)
    _title(ws, title, subtitle)
    if report.get("note"):
        ws["A3"] = f"Ghi chú: {report['note']}"
    ws["A4"] = "TỔNG KẾT KỲ"
    ws["A4"].font = Font(bold=True, color="FFFFFF")
    for c in range(1, 7):
        ws.cell(row=4, column=c).fill = HEAD_FILL
    n = len(report["rows"])
    last = FIRST_DATA_ROW + max(n, 1) - 1
    loai = f"$E${FIRST_DATA_ROW}:$E${last}"
    amt = f"$J${FIRST_DATA_ROW}:$J${last}"
    fee = f"$K${FIRST_DATA_ROW}:$K${last}"
    r = SUMMARY_ROWS
    s = report["summary"]
    items = [
        ("opening", "Số dư đầu kỳ", s["opening"]),
        ("total_in", "Tổng vào (không phí)", f'=SUMIFS({amt},{loai},"vào")+SUMIFS({amt},{loai},"Hoàn tác vào")'),
        ("total_out", "Tổng ra (không phí)", f'=SUMIFS({amt},{loai},"ra")+SUMIFS({amt},{loai},"Hoàn tác ra")'),
        ("fee_in", "Phí thu từ tiền vào", f'=SUMIFS({fee},{loai},"vào")+SUMIFS({fee},{loai},"Hoàn tác vào")'),
        ("fee_out", "Phí thu từ tiền ra", f'=SUMIFS({fee},{loai},"ra")+SUMIFS({fee},{loai},"Hoàn tác ra")'),
        ("fee_income", "Tổng phí", f"=F{r['fee_in']}+F{r['fee_out']}"),
        ("net_in", "Còn lại từ tiền vào (sau phí)", f"=F{r['total_in']}-F{r['fee_in']}"),
        ("net_out", "Ra sau phí", f"=F{r['total_out']}+F{r['fee_out']}"),
        ("closing", "Số dư cuối kỳ", f"=F{r['opening']}+F{r['total_in']}-F{r['total_out']}"),
        ("position", "Kết quả", f'=IF(F{r["closing"]}>0,"DƯƠNG",IF(F{r["closing"]}<0,"ÂM","CÂN BẰNG"))'),
        ("check", "Kiểm tra khớp sổ quỹ",
         f'=IF(F{r["closing"]}={int(ledger_closing)},"KHỚP","LỆCH - kiểm tra lại")'),
    ]
    for key, label, val in items:
        row = r[key]
        ws.cell(row=row, column=1, value=label).font = Font(bold=key in ("closing", "position"))
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
        cell = ws.cell(row=row, column=6, value=val)
        cell.number_format = NUM
        cell.alignment = Alignment(horizontal="right")
        for c in range(1, 7):
            ws.cell(row=row, column=c).border = BOX
            ws.cell(row=row, column=c).fill = SUB_FILL if key in ("closing", "position") else TOTAL_FILL
    ws.conditional_formatting.add(f"F{r['position']}", CellIsRule(operator="equal", formula=['"DƯƠNG"'], fill=GREEN, font=Font(bold=True, color="006100")))
    ws.conditional_formatting.add(f"F{r['position']}", CellIsRule(operator="equal", formula=['"ÂM"'], fill=RED, font=Font(bold=True, color="9C0006")))
    ws.conditional_formatting.add(f"F{r['check']}", CellIsRule(operator="equal", formula=['"KHỚP"'], fill=GREEN))
    ws.conditional_formatting.add(f"F{r['check']}", FormulaRule(formula=[f'LEFT(F{r["check"]},4)="LỆCH"'], fill=RED))

    ws.cell(row=FIRST_DATA_ROW - 2, column=1, value="CHI TIẾT TỪNG BILL").font = Font(bold=True, color="FFFFFF")
    for c in range(1, len(COLS) + 1):
        ws.cell(row=FIRST_DATA_ROW - 2, column=c).fill = HEAD_FILL
    for c, (head, width) in enumerate(COLS, 1):
        cell = ws.cell(row=FIRST_DATA_ROW - 1, column=c, value=head)
        cell.font = Font(bold=True)
        cell.fill = SUB_FILL
        cell.border = BOX
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = width
    ws.row_dimensions[FIRST_DATA_ROW - 1].height = 32
    for i, row in enumerate(report["rows"]):
        vals = [row["stt"], _fmt_time(row["time"]), row["bill_id"], row["entry_id"], row["loai"], row["nguoi"] or "",
                row["account"] or "", row["bank"] or "", row["content"] or "", row["amount"], row["fee"], row["net"], row["actor"]]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=FIRST_DATA_ROW + i, column=c, value=v)
            cell.border = BOX
            if c in (10, 11, 12):
                cell.number_format = NUM
            if c == 7:
                cell.number_format = "@"
    if not report["rows"]:
        ws.cell(row=FIRST_DATA_ROW, column=1, value="(chưa có giao dịch)").font = Font(italic=True, color="7F7F7F")
    ws.freeze_panes = ws.cell(row=FIRST_DATA_ROW, column=1)
    ws.auto_filter.ref = f"A{FIRST_DATA_ROW - 1}:{get_column_letter(len(COLS))}{last}"
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{FIRST_DATA_ROW - 1}:{FIRST_DATA_ROW - 1}"
    return ws


def export_xlsx(ledger, ket_id: str, path: str, include_open: bool = True) -> str:
    """Ghi toan bo lich su chot so cua mot ket ra file Excel. Tra ve duong dan file."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Tổng hợp"
    _title(ws, "BÁO CÁO CHỐT SỔ", f"Két: {ket_id}   |   Xuất lúc {_fmt_time(time.time())}")
    heads = ["Kỳ", "Ngày chốt", "Ghi chú", "Số dư đầu kỳ", "Tổng vào", "Tổng ra", "Tổng phí", "Số dư cuối kỳ", "Kết quả", "Khớp sổ"]
    widths = [14, 17, 28, 18, 18, 18, 16, 18, 14, 20]
    for c, (h, w) in enumerate(zip(heads, widths), 1):
        cell = ws.cell(row=4, column=c, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEAD_FILL
        cell.border = BOX
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = w

    periods = []
    for c in ledger.list_closings(ket_id):
        rep = ledger.report(ket_id, c["id"])
        sheet = f"Chốt {c['id']}"
        _period_sheet(wb, sheet, f"CHỐT SỔ #{c['id']}", rep, f"Két: {ket_id}   |   Chốt lúc {_fmt_time(c['closed_at'])} bởi {c['closed_by']}", c["closing_balance"])
        periods.append((sheet, f"Chốt #{c['id']}", _fmt_time(c["closed_at"]), c.get("note") or ""))
    if include_open:
        rep = ledger.report(ket_id)
        if rep["rows"]:
            _period_sheet(wb, "Kỳ đang mở", "KỲ ĐANG MỞ (chưa chốt)", rep, f"Két: {ket_id}   |   Xem trước lúc {_fmt_time(time.time())}", rep["summary"]["closing"])
            periods.append(("Kỳ đang mở", "Đang mở", "chưa chốt", "Xem trước, chưa khóa"))

    r = SUMMARY_ROWS
    for i, (sheet, label, closed_at, note) in enumerate(periods):
        row = 5 + i
        q = f"'{sheet}'!"
        vals = [label, closed_at, note, f"={q}F{r['opening']}", f"={q}F{r['total_in']}", f"={q}F{r['total_out']}",
                f"={q}F{r['fee_income']}", f"={q}F{r['closing']}", f"={q}F{r['position']}", f"={q}F{r['check']}"]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=row, column=c, value=v)
            cell.border = BOX
            if 4 <= c <= 8:
                cell.number_format = NUM
        ws.cell(row=row, column=1).hyperlink = f"#'{sheet}'!A1"
        ws.cell(row=row, column=1).font = Font(color="0563C1", underline="single")
    if periods:
        first, last = 5, 4 + len(periods)
        total_row = last + 2
        ws.cell(row=total_row, column=1, value="TỔNG CÁC KỲ").font = Font(bold=True)
        for col, letter in ((5, "E"), (6, "F"), (7, "G")):
            cell = ws.cell(row=total_row, column=col, value=f"=SUM({letter}{first}:{letter}{last})")
            cell.number_format = NUM
            cell.font = Font(bold=True)
        cell = ws.cell(row=total_row, column=8, value=f"=H{last}")
        cell.number_format = NUM
        cell.font = Font(bold=True)
        for c in range(1, len(heads) + 1):
            ws.cell(row=total_row, column=c).fill = SUB_FILL
            ws.cell(row=total_row, column=c).border = BOX
        ws.conditional_formatting.add(f"I{first}:I{last}", CellIsRule(operator="equal", formula=['"DƯƠNG"'], fill=GREEN, font=Font(bold=True, color="006100")))
        ws.conditional_formatting.add(f"I{first}:I{last}", CellIsRule(operator="equal", formula=['"ÂM"'], fill=RED, font=Font(bold=True, color="9C0006")))
        ws.conditional_formatting.add(f"J{first}:J{last}", CellIsRule(operator="equal", formula=['"KHỚP"'], fill=GREEN))
        ws.conditional_formatting.add(f"J{first}:J{last}", FormulaRule(formula=[f'LEFT(J{first},4)="LỆCH"'], fill=RED))
    else:
        ws.cell(row=5, column=1, value="(chưa có kỳ chốt nào)").font = Font(italic=True, color="7F7F7F")
    ws.freeze_panes = "A5"
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    wb.save(path)
    return path
