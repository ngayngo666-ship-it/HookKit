import os
import shutil
import subprocess
import sys

import pytest
from openpyxl import load_workbook

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
from excel_export import FIRST_DATA_ROW, SUMMARY_ROWS, export_xlsx  # noqa: E402
from ledger import Ledger  # noqa: E402
from nlu import understand  # noqa: E402

K = "nhom-ket"
needs_soffice = pytest.mark.skipif(shutil.which("soffice") is None, reason="can cai libreoffice-calc de tinh cong thuc")


def _bill(lg, amount, direction, cust, txn, content):
    r = lg.submit_bill(K, {"amount": amount, "account": "0123456789", "bank_name": "Vietcombank", "content": content, "txn": txn},
                       "2", direction, image_bytes=txn.encode())
    lg.assign_customer(r["bill_id"], cust, "2")
    return lg.confirm_bill(r["bill_id"], "2") if direction == "in" else lg.export_bill(r["bill_id"], "2")


@pytest.fixture
def scenario(tmp_path):
    lg = Ledger(str(tmp_path / "k.db"))
    lg.add_user("1", "admin")
    lg.add_user("2", "quan_tri")
    a = lg.add_customer(K, "Nguyen A", fee_in_pct=1.5, fee_out_pct=0.5, fee_out_fixed=2000, actor="2")["customer_id"]
    b = lg.add_customer(K, "Tran B", fee_in_pct=3, actor="2")["customer_id"]
    _bill(lg, 10_000_000, "in", a, "T1", "tien hang")
    _bill(lg, 4_000_000, "in", b, "T2", "dat coc")
    _bill(lg, 6_000_000, "out", a, "T3", "tra nha cung cap")
    lg.close_period(K, "2", note="Ngay 01/10")
    e = _bill(lg, 2_000_000, "in", b, "T4", "them coc")["entry_id"]
    lg.close_period(K, "2", note="Ngay 02/10")
    _bill(lg, 1_500_000, "out", a, "T5", "chi phi")
    lg.reverse_entry(e, "1")
    return lg


def _recalc(path, outdir):
    os.makedirs(outdir, exist_ok=True)
    subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", outdir, path],
                   check=True, capture_output=True, timeout=180)
    return load_workbook(os.path.join(outdir, os.path.basename(path)), data_only=True)


def test_workbook_structure_and_formulas(scenario, tmp_path):
    path = export_xlsx(scenario, K, str(tmp_path / "bc.xlsx"))
    wb = load_workbook(path)
    assert wb.sheetnames == ["Tổng hợp", "Chốt 1", "Chốt 2", "Kỳ đang mở"]
    assert all(len(n) <= 31 for n in wb.sheetnames)
    ws = wb["Chốt 1"]
    assert ws.cell(row=FIRST_DATA_ROW - 1, column=10).value == "Số tiền (không phí)"
    assert ws.cell(row=FIRST_DATA_ROW - 1, column=12).value == "Số còn lại (sau phí)"
    assert [ws.cell(row=FIRST_DATA_ROW + i, column=6).value for i in range(3)] == ["Nguyen A", "Tran B", "Nguyen A"]
    assert ws.cell(row=FIRST_DATA_ROW, column=10).value == 10_000_000
    assert ws.cell(row=FIRST_DATA_ROW, column=11).value == 150_000
    assert ws.cell(row=FIRST_DATA_ROW, column=12).value == 9_850_000
    for key in ("total_in", "total_out", "fee_in", "fee_out", "fee_income", "net_in", "net_out", "closing", "position", "check"):
        assert str(ws.cell(row=SUMMARY_ROWS[key], column=6).value).startswith("=")
    assert str(wb["Tổng hợp"]["E5"].value).startswith("='Chốt 1'!")


@needs_soffice
def test_excel_formulas_match_ledger_exactly(scenario, tmp_path):
    path = export_xlsx(scenario, K, str(tmp_path / "bc.xlsx"))
    wb = _recalc(path, str(tmp_path / "calc"))
    sheets = [(f"Chốt {c['id']}", scenario.report(K, c["id"])) for c in scenario.list_closings(K)] + [("Kỳ đang mở", scenario.report(K))]
    for name, rep in sheets:
        w, s = wb[name], rep["summary"]
        got = {k: w.cell(row=SUMMARY_ROWS[k], column=6).value for k in SUMMARY_ROWS}
        for k in ("opening", "total_in", "total_out", "fee_in", "fee_out", "fee_income", "net_in", "net_out", "closing"):
            assert got[k] == s[k], (name, k, got[k], s[k])
        assert got["position"] == s["position"] and got["check"] == "KHỚP"
    t = wb["Tổng hợp"]
    assert [t.cell(row=5, column=c).value for c in range(4, 11)] == [0, 14_000_000, 6_000_000, 302_000, 8_000_000, "DƯƠNG", "KHỚP"]
    assert [t.cell(row=9, column=c).value for c in range(5, 9)] == [14_000_000, 7_500_000, 311_500, 6_500_000]


@needs_soffice
def test_hand_edit_in_excel_is_detected(scenario, tmp_path):
    path = export_xlsx(scenario, K, str(tmp_path / "bc.xlsx"))
    wb = load_workbook(path)
    wb["Chốt 1"].cell(row=FIRST_DATA_ROW, column=10).value = 10_000_001
    wb.save(path)
    calc = _recalc(path, str(tmp_path / "calc"))
    assert calc["Chốt 1"].cell(row=SUMMARY_ROWS["check"], column=6).value.startswith("LỆCH")
    assert calc["Tổng hợp"]["J5"].value.startswith("LỆCH")


@needs_soffice
def test_negative_balance_shows_am(tmp_path):
    lg = Ledger(str(tmp_path / "n.db"), allow_negative=True)
    lg.add_user("2", "quan_tri")
    b = lg.submit_bill(K, {"amount": 700_000}, "2", "out", image_bytes=b"x")
    lg.export_bill(b["bill_id"], "2")
    lg.close_period(K, "2")
    calc = _recalc(export_xlsx(lg, K, str(tmp_path / "n.xlsx")), str(tmp_path / "calc"))
    w = calc["Chốt 1"]
    assert w.cell(row=SUMMARY_ROWS["closing"], column=6).value == -700_000
    assert w.cell(row=SUMMARY_ROWS["position"], column=6).value == "ÂM"


def test_empty_ledger_exports_without_error(tmp_path):
    lg = Ledger(str(tmp_path / "e.db"))
    wb = load_workbook(export_xlsx(lg, K, str(tmp_path / "e.xlsx")))
    assert wb.sheetnames == ["Tổng hợp"]
    assert wb["Tổng hợp"]["A5"].value == "(chưa có kỳ chốt nào)"


def test_phone_numbers_and_accounts_stay_text(scenario, tmp_path):
    ws = load_workbook(export_xlsx(scenario, K, str(tmp_path / "bc.xlsx")))["Chốt 1"]
    assert ws.cell(row=FIRST_DATA_ROW, column=7).value == "0123456789"
    assert ws.cell(row=FIRST_DATA_ROW, column=7).number_format == "@"


def test_nlu_excel_request():
    for text in ("xuất excel", "file excel", "bảng excel", "xuất bảng tính"):
        assert understand(text)["intent"] == "report", text
